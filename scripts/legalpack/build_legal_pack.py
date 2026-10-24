#!/usr/bin/env python3
"""Build the SmplWise legal-review pack (LEGALPACK-295).

Reads Markdown sources (repo docs + private planning files), writes RTL-aware .docx
(python-docx) and .pdf (headless Chrome/Edge print of an RTL HTML), a README with a
table of contents, a sha256 manifest and a ZIP.

Usage:
  py -3.12 scripts/legalpack/build_legal_pack.py --private-planning C:/cloude/smplwisebms/private/planning \
      --out private/legal-pack

Content is copied as-is (formatting only). Nothing from private-evidence/ or secrets/ is read.
"""
import argparse, hashlib, html, os, re, shutil, subprocess, sys, tempfile, zipfile
from html.parser import HTMLParser
from pathlib import Path

import markdown
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REPO = Path(__file__).resolve().parents[2]
SRC = Path(__file__).resolve().parent / "src"
PACK_NAME = "SMPLWISE_LEGAL_REVIEW_PACK"
BUILD_DATE = "2026-10-24"
FONT = "Arial"
MONO = "Consolas"

PRIV = "docs/privacy"
PLAN = "private/planning"
BRIEF = PLAN + "/legal_pack_technician_privacy_2026-10-18"

# (folder, outname, root, relpath, title_he, dir, summary_he or None, note_he or None)
# root: "repo" | "plan"
D = []


def add(folder, name, root, rel, title, direction, summary=None, note=None):
    D.append(dict(folder=folder, name=name, root=root, rel=rel, title=title, dir=direction,
                  summary=summary, note=note))


F02 = "02_privacy_org_documents"
add(F02, "00_index_C1-C13", "repo", f"{PRIV}/README.md", "אינדקס מסמכי הפרטיות של הארגון ופריטי בדיקה C1-C13 (דו-לשוני)", "rtl")
heb = [("01", "data-processing-summary", "סיכום עיבוד המידע: מה המוצר שומר, איפה, ומי מקבל"),
       ("02", "retention-and-access-summary", "סיכום שמירה וגישה: מה נמחק אוטומטית, מה ידני, ומי רואה מה"),
       ("03", "data-subject-request-handling", "שלבי טיפול בבקשות של אנשים (עיון, תיקון, מחיקה)"),
       ("04", "organisation-privacy-notice-template", "תבנית הודעת פרטיות של הארגון"),
       ("05", "database-definition-and-incident-line-template", "תבנית מסמך הגדרת מאגר ושורת נוהל אירוע"),
       ("06", "retention-policy-template", "תבנית מדיניות שמירה")]
for n, slug, t in heb:
    add(F02, f"{n}_{slug}_he", "repo", f"{PRIV}/he/{n}-{slug}.md".replace("-" + slug, "-" + slug), t + " (עברית)", "rtl")
    add(F02, f"{n}_{slug}_en", "repo", f"{PRIV}/en/{n}-{slug}.md", t + " (אנגלית, המקור)", "ltr",
        summary=f"נוסח אנגלי: המקור, והעברי ({n}) עוקב אחריו סעיף אחר סעיף. הנוסח העברי נמצא באותה תיקייה.")

F03 = "03_employee_notice"
add(F03, "01_employee-notice-draft-v3_he", "plan", f"{PLAN}/EMPLOYEE_NOTICE_DRAFT_V3_2026-10-24_HE.md",
    "טיוטה v3: הודעה לעובדים על שיתוף נתוני המכשיר עם המערכת", "rtl",
    note="טיוטת עבודה שנכתבה על ידי צוות הפיתוח. אינה מוכנה לשימוש אצל לקוח. גרסאות קודמות (v1, v2) הוחלפו ואינן בחבילה.")

F04 = "04_technician_privacy"
add(F04, "01_legal-brief-technicians_he", "plan", f"{BRIEF}/LEGAL_BRIEF_HE.md",
    "תדריך ליועץ המשפטי: הודעת פרטיות לטכנאים חיצוניים (13 שאלות בסעיף 9)", "rtl",
    note="התדריך נכתב לגרסה 2.7.1 (18.10.2026) וחלקים בו התיישנו; ראו מכתב הבקשה, סעיף 7, פריט 5.")
add(F04, "02_notice-text-proposed_he-en-ar-ru", "plan", f"{BRIEF}/NOTICE_TEXT_PROPOSED_he_en_ar_ru.md",
    "נוסחי ההודעה המוצעים לטכנאים בעברית, אנגלית, ערבית ורוסית", "rtl",
    note="הנוסחים בערבית וברוסית ממתינים לביקורת דוברי שפת אם.")
add(F04, "03_technician-notice-final_he-en-ar-ru", "repo", f"{PRIV}/technician-notice-final.md",
    "ההודעה לטכנאים: הנוסח הסופי הנוכחי בארבע שפות (ממתין לבדיקה משפטית)", "ltr",
    summary="מסמך באנגלית עם נוסחי ההודעה בעברית, ערבית ורוסית; מצב: ממתין לבדיקה משפטית.")
add(F04, "04_technician-verification-message-spec_en", "repo", f"{PRIV}/technician-verification-message-spec.md",
    "מפרט הודעת אימות הטלפון לטכנאי שנושאת את הודעת הפרטיות הקצרה", "ltr",
    summary="מפרט בלבד (אנגלית): איך המערכת מזהה טכנאי ומוסיפה לו את שורת הפרטיות בהודעת אימות המספר.")
add(F04, "05_proposal-pub03-pending-technician-retention_en", "repo", "docs/security/PROPOSAL_PUB03_PENDING_TECHNICIAN_RETENTION.md",
    "הצעה PUB-03: שמירת שם וטלפון של טכנאי שהוקלדו על ידי ספק", "ltr",
    summary="הצעה (אנגלית): שמירה ומחיקה של פרטי טכנאי שטרם אושר; האפשרות שנבחרה יושמה לפי החלטת הבעלים מ-21.10.2026.")

F05 = "05_presence_push_third_parties"
add(F05, "01_CR-027-mobile-presence-push_en", "repo", "docs/changes/CR-027-MOBILE-PRESENCE-PUSH.md",
    "מפרט CR-027: אפליקציית הנייד, רישום מכשירים, נוכחות וחיישנים, ודחיפה דרך ממסר SmplWise", "ltr",
    summary="מפרט שינוי (אנגלית): איך נוכחות עובדים וחיישני מכשיר נשלחים למערכת, מה נשמר, והדחיפה דרך ממסר של SmplWise.")
add(F05, "02_cloudflare-dpa-request-draft_en-he", "plan", f"{PLAN}/drafts_2026-10-21/cloudflare-dpa-request.md",
    "טיוטת בקשה ל-Cloudflare להסכם עיבוד מידע (DPA), טרם נשלחה", "ltr",
    summary="טיוטה באנגלית ובעברית; לא נשלחה. כוללת בראשה נקודות פתוחות לבעלים.")
add(F05, "03_whatsapp-template-meta-approval-draft_en-he", "plan", f"{PLAN}/drafts_2026-10-21/whatsapp-template-meta-approval.md",
    "טיוטת תבנית הודעת וואטסאפ לאישור Meta (קטגוריית UTILITY), טרם הוגשה", "ltr",
    summary="טיוטה באנגלית ובעברית; לא הוגשה ל-Meta.")

F06 = "06_security_background"
add(F06, "01_dependency-and-secrets-audit_en", "repo", "docs/security/DEPENDENCY_AND_SECRETS_AUDIT.md",
    "ביקורת תלויות וסודות (אנגלית)", "ltr", summary="ביקורת אבטחה של תלויות וניהול סודות במוצר (אנגלית); המקבילה העברית בקובץ 02.")
add(F06, "02_dependency-and-secrets-audit_he", "repo", "docs/security/DEPENDENCY_AND_SECRETS_AUDIT_HE.md",
    "ביקורת תלויות וסודות (עברית)", "rtl")

F07 = "07_open_items_index"
add(F07, "01_legal-pack-index-27-items_he", "plan", f"{PLAN}/LEGAL_PACK_INDEX_2026-10-24_HE.md",
    "אינדקס 27 הפריטים שדורשים בדיקה משפטית (24.10.2026)", "rtl",
    note="האינדקס מפנה גם לקבצים פנימיים שאינם בחבילה (ראו 00, סעיף 5).")

REPO_SHA = "b6eb212ec"

# ---------------------------------------------------------------------------- markdown -> html

LIST_RE = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+")


def preprocess(md: str) -> str:
    """Make loosely written Markdown parse with python-markdown (no wording change)."""
    md = md.replace("\r\n", "\n").lstrip("\ufeff")
    out, prev_blank, prev_list, in_fence = [], True, False, False
    for line in md.split("\n"):
        if line.strip().startswith("```"):
            in_fence = not in_fence
            out.append(line); prev_blank = False; prev_list = False
            continue
        if in_fence:
            out.append(line); continue
        m = LIST_RE.match(line)
        if m:
            ind = len(m.group(1).replace("\t", "    "))
            if ind in (2, 3): ind = 4
            elif ind in (5, 6, 7): ind = 8
            line = " " * ind + line.lstrip()
            if not prev_blank and not prev_list:
                out.append("")
            prev_list = True
        else:
            if line.strip() and prev_list and not line.startswith((" ", "\t")) and not prev_blank:
                out.append("")
            prev_list = prev_list and (line.strip() == "" or line.startswith((" ", "\t")))
        # table must be preceded by a blank line
        if line.lstrip().startswith("|") and out and out[-1].strip() and not out[-1].lstrip().startswith("|"):
            out.append("")
        out.append(line)
        prev_blank = line.strip() == ""
    return "\n".join(out)


def md_to_html(md: str) -> str:
    return markdown.markdown(preprocess(md), extensions=["tables", "fenced_code", "sane_lists"])


# ---------------------------------------------------------------------------- tiny DOM

VOID = {"br", "hr", "img"}


class Node:
    def __init__(self, tag, attrs=None):
        self.tag, self.attrs, self.kids = tag, dict(attrs or {}), []

    def text(self):
        return "".join(k if isinstance(k, str) else k.text() if k.tag != "br" else "\n" for k in self.kids)


class Builder(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node("root"); self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        n = Node(tag, attrs); self.stack[-1].kids.append(n)
        if tag not in VOID: self.stack.append(n)

    def handle_endtag(self, tag):
        if tag in VOID: return
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]; break

    def handle_data(self, data):
        self.stack[-1].kids.append(data)


def parse_html(h):
    b = Builder(); b.feed(h); return b.root


# ---------------------------------------------------------------------------- direction

RTL_RE = re.compile(r"[\u0590-\u08FF\uFB1D-\uFDFF\uFE70-\uFEFF]")
LTR_RE = re.compile(r"[A-Za-z\u00C0-\u024F\u0400-\u04FF]")


def is_rtl(text, default):
    r = len(RTL_RE.findall(text)); l = len(LTR_RE.findall(text))
    if r == 0 and l == 0: return default
    first = None
    for ch in text:
        if RTL_RE.match(ch): first = True; break
        if LTR_RE.match(ch): first = False; break
    # first strong character wins unless the other script clearly dominates
    if first is True and r * 3 >= l: return True
    if first is False and l * 3 >= r: return False
    return r >= l


# ---------------------------------------------------------------------------- docx helpers

def set_run_font(run, size=None, bold=None, italic=None, mono=False, rtl=False, color=None):
    rpr = run._r.get_or_add_rPr()
    f = rpr.find(qn("w:rFonts"))
    if f is None:
        f = OxmlElement("w:rFonts"); rpr.insert(0, f)
    name = MONO if mono else FONT
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        f.set(qn(a), name)
    if size:
        run.font.size = Pt(size)
        szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), str(int(size * 2))); rpr.append(szcs)
    if bold:
        run.bold = True
        b = OxmlElement("w:bCs"); rpr.append(b)
    if italic:
        run.italic = True
        i = OxmlElement("w:iCs"); rpr.append(i)
    if color: run.font.color.rgb = RGBColor.from_string(color)
    if rtl:
        rtl_el = OxmlElement("w:rtl"); rpr.append(rtl_el)


def shade(el_pr, fill):
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
    el_pr.append(sh)


def new_par(container, rtl, style=None):
    p = container.add_paragraph(style=style) if style else container.add_paragraph()
    if rtl:
        p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
    return p


def collect_runs(node, bold=False, italic=False, code=False, out=None):
    out = [] if out is None else out
    for k in node.kids:
        if isinstance(k, str):
            out.append((k, bold, italic, code))
        elif k.tag == "br":
            out.append(("\n", bold, italic, code))
        elif k.tag in ("strong", "b"):
            collect_runs(k, True, italic, code, out)
        elif k.tag in ("em", "i"):
            collect_runs(k, bold, True, code, out)
        elif k.tag == "code":
            collect_runs(k, bold, italic, True, out)
        elif k.tag in ("ul", "ol", "table"):
            continue
        else:
            collect_runs(k, bold, italic, code, out)
    return out


def fill_par(p, runs, rtl, size, base_bold=False, color=None):
    for text, b, i, code in runs:
        parts = text.split("\n")
        for idx, part in enumerate(parts):
            if idx > 0: p.add_run().add_break()
            if not part: continue
            r = p.add_run(part)
            set_run_font(r, size=size, bold=b or base_bold, italic=i, mono=code,
                         rtl=(rtl and not code), color=color)
            if code:
                shade(r._r.get_or_add_rPr(), "F0F0F0")


def plain(runs): return "".join(t for t, *_ in runs)



ORDER = {
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath".split(),
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange".split(),
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook".split(),
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark".split(),
}


def normalize(root):
    """Re-order property children into the OOXML schema sequence (Word is strict about it)."""
    for tag, order in ORDER.items():
        for el in root.iter(qn("w:" + tag)):
            kids = list(el)
            seen, uniq = set(), []
            for k in kids:
                name = k.tag.split("}")[1]
                if name in seen and name in ("rtl", "bidi", "bidiVisual", "szCs", "bCs", "iCs", "shd", "lang"):
                    el.remove(k); continue
                seen.add(name); uniq.append(k)
            uniq.sort(key=lambda k: order.index(k.tag.split("}")[1]) if k.tag.split("}")[1] in order else len(order))
            for k in uniq: el.remove(k)
            for k in uniq: el.append(k)


class DocxWriter:
    def __init__(self, default_rtl):
        self.doc = Document()
        self.default_rtl = default_rtl
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)
        st = self.doc.styles
        n = st["Normal"]; n.font.name = FONT; n.font.size = Pt(11)
        rpr = n.element.get_or_add_rPr()
        f = rpr.find(qn("w:rFonts"))
        if f is None:
            f = OxmlElement("w:rFonts"); rpr.insert(0, f)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): f.set(qn(a), FONT)
        lang = OxmlElement("w:lang"); lang.set(qn("w:val"), "en-US"); lang.set(qn("w:bidi"), "he-IL"); rpr.append(lang)
        n.paragraph_format.space_after = Pt(6); n.paragraph_format.line_spacing = 1.15
        for lvl, sz in ((1, 18), (2, 15), (3, 13), (4, 12)):
            h = st[f"Heading {lvl}"]; h.font.name = FONT; h.font.size = Pt(sz); h.font.bold = True
            h.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
            hr = h.element.get_or_add_rPr(); hf = hr.find(qn("w:rFonts"))
            if hf is None:
                hf = OxmlElement("w:rFonts"); hr.insert(0, hf)
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): hf.set(qn(a), FONT)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                if hf.get(qn(a)) is not None: del hf.attrib[qn(a)]
            h.paragraph_format.space_before = Pt(14 if lvl <= 2 else 10); h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
        self.footer(sec)

    def footer(self, sec):
        p = sec.footer.paragraphs[0]
        p.alignment = 1
        r = p.add_run("SmplWise Arx - Legal Review Pack - DRAFT for counsel - Page ")
        set_run_font(r, 8, color="666666")
        for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            run = p.add_run(); set_run_font(run, 8, color="666666")
            if kind:
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); run._r.append(fc)
            else:
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; run._r.append(it)

    # ---- blocks
    def render(self, root, container=None, ctx=None):
        container = container or self.doc
        ctx = ctx or {"level": 0, "size": 11}
        for k in root.kids:
            if isinstance(k, str):
                if k.strip():
                    p = new_par(container, self.default_rtl); fill_par(p, [(k.strip(), False, False, False)], self.default_rtl, ctx["size"])
                continue
            t = k.tag
            if t in ("h1", "h2", "h3", "h4", "h5", "h6"):
                runs = collect_runs(k); rtl = is_rtl(plain(runs), self.default_rtl)
                lvl = min(int(t[1]), 4)
                p = new_par(container, rtl, style=f"Heading {lvl}")
                fill_par(p, runs, rtl, None, base_bold=True, color="1F2A44")
            elif t == "p":
                self.para(container, k, ctx)
            elif t in ("ul", "ol"):
                self.list(container, k, ctx)
            elif t == "table":
                self.table(container, k, ctx)
            elif t == "pre":
                for line in k.text().rstrip("\n").split("\n"):
                    p = new_par(container, False)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.left_indent = Cm(0.5)
                    if line:
                        r = p.add_run(line); set_run_font(r, 9, mono=True)
                    shade(p._p.get_or_add_pPr(), "F4F4F4")
                self.gap(container)
            elif t == "blockquote":
                sub = dict(ctx); sub["quote"] = True
                self.render(k, container, sub)
            elif t == "hr":
                p = new_par(container, False)
                pb = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
                for a, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "999999")): b.set(qn("w:" + a), v)
                pb.append(b); p._p.get_or_add_pPr().append(pb)
            else:
                self.render(k, container, ctx)

    def gap(self, container):
        p = new_par(container, False); p.paragraph_format.space_after = Pt(4)

    def para(self, container, k, ctx, extra_indent=None):
        runs = collect_runs(k)
        if not plain(runs).strip(): return
        rtl = is_rtl(plain(runs), self.default_rtl)
        p = new_par(container, rtl)
        if ctx.get("quote"):
            p.paragraph_format.left_indent = Cm(0.8)
        fill_par(p, runs, rtl, ctx["size"])
        return p

    def list(self, container, k, ctx):
        ordered = k.tag == "ol"; n = int(k.attrs.get("start", 1)) if ordered else 1
        for li in [x for x in k.kids if not isinstance(x, str) and x.tag == "li"]:
            runs = collect_runs(li)
            rtl = is_rtl(plain(runs), self.default_rtl)
            lvl = ctx["level"]
            p = new_par(container, rtl)
            p.paragraph_format.left_indent = Cm(0.9 + 0.8 * lvl); p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(3)
            marker = (f"{n}." if ordered else ("•" if lvl == 0 else "–")) + "\t"
            fill_par(p, [(marker, False, False, False)] + runs, rtl, ctx["size"])
            n += 1
            for sub in [x for x in li.kids if not isinstance(x, str) and x.tag in ("ul", "ol")]:
                c2 = dict(ctx); c2["level"] = lvl + 1
                self.list(container, sub, c2)
            for sub in [x for x in li.kids if not isinstance(x, str) and x.tag == "table"]:
                self.table(container, sub, ctx)

    def table(self, container, k, ctx):
        rows = []
        for sec in [x for x in k.kids if not isinstance(x, str)]:
            if sec.tag in ("thead", "tbody"):
                for tr in [r for r in sec.kids if not isinstance(r, str) and r.tag == "tr"]:
                    rows.append((sec.tag == "thead", tr))
            elif sec.tag == "tr":
                rows.append((False, sec))
        if not rows: return
        ncols = max(len([c for c in tr.kids if not isinstance(c, str)]) for _, tr in rows)
        # table direction from header text
        sample = " ".join(c.text() for _, tr in rows[:2] for c in tr.kids if not isinstance(c, str))
        trtl = is_rtl(sample, self.default_rtl)
        tbl = self.doc.add_table(rows=len(rows), cols=ncols) if container is self.doc else container.add_table(rows=len(rows), cols=ncols)
        tbl.style = "Table Grid"
        tpr = tbl._tbl.tblPr
        if trtl: tpr.append(OxmlElement("w:bidiVisual"))
        size = 9.5 if ncols <= 4 else 8.5 if ncols <= 6 else 7.5
        # explicit equal-ish widths so Word does not squash columns
        total = 17.0
        for ri, (hdr, tr) in enumerate(rows):
            cells = [c for c in tr.kids if not isinstance(c, str)]
            for ci in range(ncols):
                cell = tbl.cell(ri, ci)
                cell.width = Cm(total / ncols)
                if ci >= len(cells): continue
                runs = collect_runs(cells[ci])
                crtl = is_rtl(plain(runs), trtl)
                p = cell.paragraphs[0]
                if crtl: p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
                p.paragraph_format.space_after = Pt(2)
                fill_par(p, runs, crtl, size, base_bold=(hdr or cells[ci].tag == "th"))
                if hdr or cells[ci].tag == "th":
                    shade(cell._tc.get_or_add_tcPr(), "E3E8F2")
            if hdr:
                trpr = tbl.rows[ri]._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:tblHeader"))
        self.gap(container)

    def cover(self, title, meta, summary, direction_rtl):
        t = new_par(self.doc, True, style="Title")
        fill_par(t, [(title, True, False, False)], True, 17, color="1F2A44")
        tbl = self.doc.add_table(rows=len(meta), cols=2); tbl.style = "Table Grid"
        tbl._tbl.tblPr.append(OxmlElement("w:bidiVisual"))
        for i, (k, v) in enumerate(meta):
            for j, txt in enumerate((k, v)):
                c = tbl.cell(i, j); c.width = Cm(4 if j == 0 else 13)
                p = c.paragraphs[0]
                rtl = is_rtl(txt, True)
                if rtl: p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
                fill_par(p, [(txt, j == 0, False, False)], rtl, 9.5)
                if j == 0: shade(c._tc.get_or_add_tcPr(), "E3E8F2")
        if summary:
            p = new_par(self.doc, True); p.paragraph_format.space_before = Pt(6)
            fill_par(p, [(summary, True, False, False)], True, 10.5)
        sep = new_par(self.doc, False)
        pb = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
        for a, v in (("val", "single"), ("sz", "8"), ("space", "1"), ("color", "1F2A44")): b.set(qn("w:" + a), v)
        pb.append(b); sep._p.get_or_add_pPr().append(pb)

    def save(self, path):
        normalize(self.doc.element)
        normalize(self.doc.styles.element)
        for sec in self.doc.sections: normalize(sec.footer._element)
        self.doc.save(str(path))


# ---------------------------------------------------------------------------- html / pdf

CSS = """
@page { size: A4; margin: 18mm 16mm 20mm 16mm;
  @bottom-center { content: "SmplWise Arx - Legal Review Pack - DRAFT for counsel - Page " counter(page) " / " counter(pages); direction:ltr; font: 8pt Arial; color:#666; } }
html { font-family: Arial, 'Segoe UI', Tahoma, sans-serif; font-size: 10.5pt; line-height: 1.5; color:#111; }
body { margin:0; }
h1,h2,h3,h4,h5,h6 { color:#1f2a44; line-height:1.25; page-break-after: avoid; margin: 1.1em 0 .4em; }
h1 { font-size: 19pt; } h2 { font-size: 15pt; border-bottom:1px solid #c9d0de; padding-bottom:2px; } h3 { font-size: 13pt; } h4 { font-size: 11.5pt; }
p { margin: .4em 0; } ul, ol { padding-inline-start: 1.8em; margin: .4em 0; } li { margin: .15em 0; }
table { border-collapse: collapse; width:100%; margin:.6em 0; page-break-inside:auto; }
tr { page-break-inside: avoid; }
th, td { border:1px solid #9aa3b5; padding:3px 6px; vertical-align: top; font-size: 9pt; overflow-wrap:anywhere; }
th { background:#e3e8f2; }
code { font-family: Consolas, 'Courier New', monospace; font-size: 9pt; background:#f0f0f0; padding:0 2px; direction:ltr; unicode-bidi: embed; overflow-wrap:anywhere; }
pre { direction:ltr; text-align:left; background:#f4f4f4; padding:6px; font-size:8.5pt; white-space:pre-wrap; overflow-wrap:anywhere; }
pre code { background:none; }
blockquote { margin: .5em 1em; padding-inline-start: .8em; border-inline-start: 3px solid #c9d0de; color:#333; }
hr { border:0; border-top:1px solid #999; margin: 1em 0; }
.cover { margin-bottom: 14px; }
.cover h1.title { font-size: 17pt; margin-top:0; }
.cover table th { width: 22%; }
.cover .summary { font-weight: bold; margin: 8px 0; }
.cover { border-bottom: 2px solid #1f2a44; padding-bottom: 8px; }
"""


def add_dir_auto(h):
    # lists get an explicit direction from their first item (dir=auto on ul/ol loses the markers in Chrome print)
    def lst(m):
        txt = re.sub(r"<[^>]+>", "", m.group(2))
        d = "rtl" if is_rtl(html.unescape(txt), True) else "ltr"
        return f'<{m.group(1)} dir="{d}">{m.group(2)}'
    h = re.sub(r"<(ul|ol)>(\s*<li>.*?)(?=</li>)", lst, h, flags=re.S)
    return re.sub(r"<(p|li|h[1-6]|td|th|blockquote)(?=[\s>])", r'<\1 dir="auto"', h)


def make_html(title, body_html, meta, summary, direction, show_cover=True):
    rows = "".join(f"<tr><th dir='rtl'>{html.escape(k)}</th><td dir='auto'>{html.escape(v)}</td></tr>" for k, v in meta)
    cover = (f"<div class='cover' dir='rtl'><h1 class='title'>{html.escape(title)}</h1><table dir='rtl'>{rows}</table>"
             + (f"<p class='summary'>{html.escape(summary)}</p>" if summary else "") + "</div>")
    return (f"<!doctype html><html lang='he' dir='{direction}'><head><meta charset='utf-8'><title>{html.escape(title)}</title>"
            f"<style>{CSS}</style></head><body dir='{direction}'>{cover if show_cover else ''}{add_dir_auto(body_html)}</body></html>")


def find_browser():
    for p in (r"C:\Program Files\Google\Chrome\Application\chrome.exe",
              r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
              r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
              r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"):
        if os.path.exists(p): return p
    return None


def html_to_pdf(browser, html_path, pdf_path, profile):
    cmd = [browser, "--headless=new", "--disable-gpu", "--no-sandbox", "--no-pdf-header-footer",
           f"--user-data-dir={profile}", f"--print-to-pdf={Path(pdf_path).resolve()}", Path(html_path).resolve().as_uri()]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
    if not Path(pdf_path).exists() or Path(pdf_path).stat().st_size < 1000:
        raise RuntimeError(f"pdf failed for {html_path}: {r.stderr[-300:]}")


# ---------------------------------------------------------------------------- scan

SECRET_PATTERNS = [
    (re.compile(r"\b(?:\d{1,3}\.){3}\d{1,3}\b"), "ipv4"),
    (re.compile(r"\b(?:[0-9A-Fa-f]{2}[:-]){5}[0-9A-Fa-f]{2}\b"), "mac"),
    (re.compile(r"(?i)\b(bearer\s+[A-Za-z0-9._-]{16,}|sk-[A-Za-z0-9]{16,}|eyJ[A-Za-z0-9_-]{20,})"), "token"),
    (re.compile(r"(?i)(password|passwd|secret|api[_-]?key)\s*[:=]\s*\S{6,}"), "credential"),
    (re.compile(r"(?i)\bserial\s*(no|number|#)?\s*[:=]\s*\S+"), "serial"),
]
SCAN_ALLOW = {"172.30.32.0"}


def scan(name, text):
    hits = []
    for rx, kind in SECRET_PATTERNS:
        for m in rx.finditer(text):
            if m.group(0) in SCAN_ALLOW: continue
            hits.append((kind, m.group(0)[:40]))
    return hits


# ---------------------------------------------------------------------------- main

def sha256(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for ch in iter(lambda: f.read(1 << 20), b""): h.update(ch)
    return h.hexdigest()


def build_one(d, plan_root, out_root, browser, profile, tmp):
    if d["root"] == "repo": src = REPO / d["rel"]
    else: src = Path(plan_root) / Path(d["rel"]).relative_to(PLAN)
    text = src.read_text(encoding="utf-8")
    hits = scan(d["name"], text)
    if hits: raise SystemExit(f"SCAN HIT in {src}: {hits}")
    return render_doc(d["title"], text, d["dir"], d.get("summary"), d.get("note"),
                      source_path=d["rel"], root=d["root"], outdir=out_root / d["folder"], name=d["name"],
                      browser=browser, profile=profile, tmp=tmp)


def render_doc(title, md, direction, summary, note, source_path, root, outdir, name, browser, profile, tmp, authored=False):
    outdir.mkdir(parents=True, exist_ok=True)
    body_html = md_to_html(md)
    if authored:
        meta = [("מסמך", title), ("הוכן", BUILD_DATE + " (נכתב עבור חבילה זו)"),
                ("סטטוס", "טיוטה לבדיקת עורך דין. אין בו ייעוץ משפטי.")]
    else:
        ver = (f"origin/main @ {REPO_SHA}" if root == "repo" else "קובץ עבודה בתיקיית התכנון (אינו מנוהל בגרסאות במאגר)")
        meta = [("מסמך", title), ("נתיב המקור", source_path), ("גרסת המקור", ver), ("הוכן לחבילה", BUILD_DATE),
                ("מצב", "הנוסח המקורי נשמר ללא שינוי; עיצוב בלבד. טיוטה לבדיקת עורך דין, אינו ייעוץ משפטי.")]
        if note: meta.append(("הערה", note))
    rtl = direction == "rtl"
    w = DocxWriter(rtl)
    if not authored:
        w.cover(title, meta, summary, rtl)
    w.render(parse_html(body_html))
    docx_path = outdir / f"{name}.docx"
    w.save(docx_path)
    html_path = Path(tmp) / f"{name}.html"
    html_path.write_text(make_html(title, body_html, meta, summary, direction, not authored), encoding="utf-8")
    pdf_path = outdir / f"{name}.pdf"
    if browser:
        html_to_pdf(browser, html_path, pdf_path, profile)
    else:
        shutil.copy(html_path, outdir / f"{name}.print.html")
    return docx_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--private-planning", required=True)
    ap.add_argument("--out", default=str(REPO / "private" / "legal-pack"))
    a = ap.parse_args()
    out = Path(a.out).resolve(); pack = out / PACK_NAME
    if pack.exists(): shutil.rmtree(pack)
    pack.mkdir(parents=True)
    browser = find_browser()
    print("browser:", browser)
    tmp = tempfile.mkdtemp(prefix="legalpack_"); profile = tempfile.mkdtemp(prefix="legalpack_prof_")
    try:
        # table of contents for README
        toc = ["| תיקייה / קובץ | מסמך | שפה | מקור |", "|---|---|---|---|"]
        toc.append("| `00_README_INSTRUCTIONS_HE` | הוראות קריאה (קובץ זה) | עברית | נכתב לחבילה |")
        toc.append("| `01_REQUEST_TO_COUNSEL_HE` | מכתב הבקשה לעורך הדין | עברית | נכתב לחבילה |")
        last = None
        for d in D:
            lang = {"rtl": "עברית", "ltr": "אנגלית / רב-לשוני"}[d["dir"]]
            if d["name"].endswith("_he") and d["dir"] == "rtl": lang = "עברית"
            if "_he-en" in d["name"] or "en-he" in d["name"]: lang = "רב-לשוני"
            toc.append(f"| `{d['folder']}/{d['name']}` | {d['title']} | {lang} | `{d['rel']}` |")
        readme = (SRC / "00_README_INSTRUCTIONS_HE.md").read_text(encoding="utf-8").replace("{{TOC}}", "\n".join(toc))
        render_doc("00 הוראות קריאה לחבילת הבדיקה המשפטית", readme, "rtl", None, None, "", "authored", pack,
                   "00_README_INSTRUCTIONS_HE", browser, profile, tmp, authored=True)
        letter = (SRC / "01_REQUEST_TO_COUNSEL_HE.md").read_text(encoding="utf-8")
        for rx, kind in SECRET_PATTERNS:
            assert not rx.search(letter + readme), kind
        render_doc("01 בקשה לעורך הדין: סקירה משפטית של מסמכי SmplWise Arx", letter, "rtl", None, None, "", "authored", pack,
                   "01_REQUEST_TO_COUNSEL_HE", browser, profile, tmp, authored=True)
        for d in D:
            print("build", d["folder"], d["name"])
            build_one(d, a.private_planning, pack, browser, profile, tmp)
    finally:
        shutil.rmtree(tmp, ignore_errors=True); shutil.rmtree(profile, ignore_errors=True)

    # manifest + zip
    lines = []
    for p in sorted(pack.rglob("*")):
        if p.is_file() and p.name != "MANIFEST_SHA256.txt":
            lines.append(f"{sha256(p)}  {p.relative_to(pack).as_posix()}")
    (pack / "MANIFEST_SHA256.txt").write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")
    zpath = out / f"{PACK_NAME}.zip"
    if zpath.exists(): zpath.unlink()
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for p in sorted(pack.rglob("*")):
            if p.is_file(): z.write(p, f"{PACK_NAME}/{p.relative_to(pack).as_posix()}")
    print("zip", zpath, zpath.stat().st_size, sha256(zpath))


if __name__ == "__main__":
    main()
